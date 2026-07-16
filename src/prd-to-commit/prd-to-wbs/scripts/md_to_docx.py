#!/usr/bin/env python3
"""
将Markdown转换为Word文档。
"""

import sys
import argparse
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    print("请安装python-docx: pip install python-docx")
    sys.exit(1)

def parse_markdown_line(line, doc):
    """解析Markdown行并添加到Word文档"""
    line = line.rstrip()
    if not line:
        doc.add_paragraph()
        return

    # 标题
    if line.startswith('# '):
        p = doc.add_heading(line[2:], level=1)
    elif line.startswith('## '):
        p = doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        p = doc.add_heading(line[4:], level=3)
    elif line.startswith('#### '):
        p = doc.add_heading(line[5:], level=4)
    elif line.startswith('> '):
        p = doc.add_paragraph(line[2:])
        p.paragraph_format.left_indent = Inches(0.25)
        p.style = doc.styles['Intense Quote']
    elif line.startswith('- ') or line.startswith('* '):
        # 处理列表
        p = doc.add_paragraph(line[2:], style='List Bullet')
    elif '|' in line and line.startswith('|'):  # 表格
        # 简单处理表格行，这里不实现完整表格解析
        p = doc.add_paragraph(line)
    else:
        # 普通段落
        p = doc.add_paragraph(line)

    # 处理内联样式（简化版本）
    if '**' in line:
        # 粗体
        pass  # python-docx的内联样式处理较复杂，这里省略

def md_to_docx(md_path, docx_path):
    """转换Markdown文件到Word文档"""
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)

    # 解析每一行
    for line in lines:
        parse_markdown_line(line, doc)

    doc.save(docx_path)

def main():
    parser = argparse.ArgumentParser(description='Markdown转Word文档')
    parser.add_argument('input', help='输入Markdown文件')
    parser.add_argument('output', help='输出Word文件')

    args = parser.parse_args()

    if not Path(args.input).exists():
        print(f"文件不存在: {args.input}")
        sys.exit(1)

    try:
        md_to_docx(args.input, args.output)
        print(f"转换完成: {args.output}")
    except Exception as e:
        print(f"转换失败: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()