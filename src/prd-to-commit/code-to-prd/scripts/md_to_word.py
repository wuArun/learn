#!/usr/bin/env python3
"""
Convert Markdown PRD to Word document with proper formatting.
"""

import sys
import re
from pathlib import Path

def convert_md_to_docx(md_file, docx_file):
    """Convert markdown to Word document."""
    
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        print("Error: python-docx not installed. Run: pip install python-docx")
        sys.exit(1)
    
    # Read markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Create Word document
    doc = Document()
    
    # Set default font for document
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    # Configure heading styles
    for i in range(1, 10):
        try:
            heading_style = doc.styles[f'Heading {i}']
            heading_style.font.name = 'Microsoft YaHei'
            heading_style.font.bold = True
            heading_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            if i == 1:
                heading_style.font.size = Pt(18)
            elif i == 2:
                heading_style.font.size = Pt(16)
            elif i == 3:
                heading_style.font.size = Pt(14)
            else:
                heading_style.font.size = Pt(12)
        except:
            pass
    
    # Parse markdown line by line
    lines = md_content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # Handle headers
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_heading(text, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Apply font
            for run in p.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.bold = True
                run.font.size = Pt(20)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        elif line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_heading(text, level=1)
            for run in p.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.bold = True
                run.font.size = Pt(16)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        elif line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_heading(text, level=2)
            for run in p.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.bold = True
                run.font.size = Pt(14)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        elif line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_heading(text, level=3)
            for run in p.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.bold = True
                run.font.size = Pt(12)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        elif line.startswith('##### '):
            text = line[6:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Microsoft YaHei'
            run.font.bold = True
            run.font.size = Pt(11)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        # Handle tables
        elif line.startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            i -= 1  # Back up one line
            
            if len(table_lines) >= 2:
                # Parse header
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                
                # Create table
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Light Grid Accent 1'
                
                # Add header
                hdr_cells = table.rows[0].cells
                for j, header in enumerate(headers):
                    hdr_cells[j].text = header
                    # Format header cell
                    for paragraph in hdr_cells[j].paragraphs:
                        for run in paragraph.runs:
                            run.font.name = 'Microsoft YaHei'
                            run.font.bold = True
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                
                # Add rows
                for row_line in table_lines[2:]:
                    cells = [c.strip() for c in row_line.split('|')[1:-1]]
                    if cells and any(cells):
                        row_cells = table.add_row().cells
                        for j, cell in enumerate(cells):
                            if j < len(row_cells):
                                # Process bold text in cells
                                process_cell_text(row_cells[j], cell)
        
        # Handle bullet lists
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            process_paragraph_text(p, text)
        
        # Handle numbered lists
        elif re.match(r'^\d+\.\s', line.strip()):
            text = re.sub(r'^\d+\.\s', '', line.strip())
            p = doc.add_paragraph(style='List Number')
            process_paragraph_text(p, text)
        
        # Handle regular paragraphs
        elif line.strip():
            p = doc.add_paragraph()
            process_paragraph_text(p, line)
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"Word document saved: {docx_file}")


def process_paragraph_text(paragraph, text):
    """Process inline formatting for paragraph."""
    from docx.oxml.ns import qn
    
    # Pattern for bold text: **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    
    for part in parts:
        if not part:
            continue
        
        run = paragraph.add_run()
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            run.text = part[2:-2]
            run.font.bold = True
        else:
            # Regular text
            run.text = part


def process_cell_text(cell, text):
    """Process inline formatting for table cell."""
    from docx.oxml.ns import qn
    
    # Clear cell and get paragraph
    cell.text = ''
    p = cell.paragraphs[0]
    
    # Pattern for bold text: **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    
    for part in parts:
        if not part:
            continue
        
        run = p.add_run()
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            run.text = part[2:-2]
            run.font.bold = True
        else:
            # Regular text
            run.text = part


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    
    md_file = sys.argv[1]
    docx_file = sys.argv[2]
    
    convert_md_to_docx(md_file, docx_file)