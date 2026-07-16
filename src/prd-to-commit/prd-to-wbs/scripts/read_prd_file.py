#!/usr/bin/env python3
"""
读取各种格式的PRD文件，提取文本内容。
支持格式：.txt, .md, .docx, .pptx, .xlsx
"""

import sys
import os
from pathlib import Path

def read_txt_md(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        return f.read()

def read_docx(filepath):
    from docx import Document
    doc = Document(filepath)
    return '\n'.join([para.text for para in doc.paragraphs])

def read_pptx(filepath):
    from pptx import Presentation
    prs = Presentation(filepath)
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(shape.text)
    return '\n'.join(texts)

def read_xlsx(filepath):
    from openpyxl import load_workbook
    wb = load_workbook(filepath, data_only=True)
    texts = []
    for sheet in wb.worksheets:
        texts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            row_text = ' | '.join([str(cell) for cell in row if cell])
            if row_text.strip():
                texts.append(row_text)
    return '\n'.join(texts)

def main():
    if len(sys.argv) < 2:
        print("Usage: python read_prd_file.py <filepath>")
        sys.exit(1)

    filepath = sys.argv[1]
    ext = Path(filepath).suffix.lower()

    readers = {
        '.txt': read_txt_md,
        '.md': read_txt_md,
        '.docx': read_docx,
        '.pptx': read_pptx,
        '.xlsx': read_xlsx,
    }

    if ext not in readers:
        print(f"Unsupported file format: {ext}")
        print(f"Supported formats: {', '.join(readers.keys())}")
        sys.exit(1)

    try:
        content = readers[ext](filepath)
        print(content)
    except Exception as e:
        print(f"Error reading file: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()